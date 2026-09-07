from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException, UploadFile
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.ai import ExamRubric
from app.models.exam import MockExam, PracticeExam

MAX_RUBRIC_BYTES = 10 * 1024 * 1024
ALLOWED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def normalize_exam_type(exam_type: str) -> str:
    value = exam_type.strip().lower()
    if value not in {"practice", "mock"}:
        raise HTTPException(422, "exam_type must be practice or mock")
    return value.title()


def _exam_filter(exam_type: str, exam_id: int):
    return (PracticeExam, PracticeExam.id == exam_id) if exam_type == "Practice" else (MockExam, MockExam.id == exam_id)


def owned_exam(db: Session, exam_type: str, exam_id: int, teacher_id: int):
    model, condition = _exam_filter(exam_type, exam_id)
    exam = db.query(model).filter(condition, model.teacher_id == teacher_id).first()
    if exam is None:
        raise HTTPException(404, "Exam not found")
    return exam


def _rubric_exam_filter(exam_type: str, exam_id: int):
    return ExamRubric.practice_exam_id == exam_id if exam_type == "Practice" else ExamRubric.mock_exam_id == exam_id


def latest_rubric(db: Session, exam_type: str, exam_id: int, teacher_id: int) -> ExamRubric | None:
    return db.query(ExamRubric).filter(
        ExamRubric.teacher_id == teacher_id,
        ExamRubric.exam_type == exam_type,
        _rubric_exam_filter(exam_type, exam_id),
    ).order_by(ExamRubric.version.desc(), ExamRubric.id.desc()).first()


def _extract_text(data: bytes, suffix: str) -> str:
    if suffix in {".txt", ".md", ".markdown"}:
        return data.decode("utf-8-sig")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise HTTPException(503, "PDF extraction dependency is not installed") from exc
        reader = PdfReader(BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise HTTPException(503, "DOCX extraction dependency is not installed") from exc
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)
    raise HTTPException(422, "Supported rubric formats: PDF, DOCX, TXT, Markdown")


def _safe_suffix(filename: str | None) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise HTTPException(422, "Supported rubric formats: PDF, DOCX, TXT, Markdown")
    return suffix


def _save_file(data: bytes, suffix: str) -> str:
    directory = Path(settings.UPLOAD_DIR) / "rubrics"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{uuid4()}{suffix}"
    path.write_bytes(data)
    return str(path)


def upload_rubric(db: Session, exam_type: str, exam_id: int, teacher_id: int, file: UploadFile) -> ExamRubric:
    exam_type = normalize_exam_type(exam_type)
    owned_exam(db, exam_type, exam_id, teacher_id)
    suffix = _safe_suffix(file.filename)
    data = file.file.read(MAX_RUBRIC_BYTES + 1)
    if len(data) > MAX_RUBRIC_BYTES:
        raise HTTPException(413, "Rubric file must not exceed 10 MB")
    try:
        content_text = _extract_text(data, suffix).strip()
    except UnicodeDecodeError as exc:
        raise HTTPException(422, "Text rubric must be UTF-8 encoded") from exc
    if not content_text:
        raise HTTPException(422, "Rubric file contains no readable text")
    current = latest_rubric(db, exam_type, exam_id, teacher_id)
    path = _save_file(data, suffix)
    rubric = ExamRubric(
        teacher_id=teacher_id,
        exam_type=exam_type,
        practice_exam_id=exam_id if exam_type == "Practice" else None,
        mock_exam_id=exam_id if exam_type == "Mock" else None,
        content_text=content_text,
        original_filename=file.filename,
        mime_type=file.content_type,
        storage_path=path,
        version=(current.version + 1) if current else 1,
    )
    db.add(rubric)
    db.flush()
    db.refresh(rubric)
    return rubric


def update_rubric_text(db: Session, exam_type: str, exam_id: int, teacher_id: int, content_text: str) -> ExamRubric:
    exam_type = normalize_exam_type(exam_type)
    owned_exam(db, exam_type, exam_id, teacher_id)
    text = content_text.strip()
    if not text:
        raise HTTPException(422, "Rubric text cannot be empty")
    current = latest_rubric(db, exam_type, exam_id, teacher_id)
    if current is None:
        raise HTTPException(404, "Rubric not found")
    rubric = ExamRubric(
        teacher_id=teacher_id,
        exam_type=exam_type,
        practice_exam_id=exam_id if exam_type == "Practice" else None,
        mock_exam_id=exam_id if exam_type == "Mock" else None,
        content_text=text,
        original_filename=current.original_filename,
        mime_type="text/plain",
        storage_path=current.storage_path,
        version=current.version + 1,
    )
    db.add(rubric)
    db.flush()
    db.refresh(rubric)
    return rubric


def delete_rubric(db: Session, exam_type: str, exam_id: int, teacher_id: int) -> None:
    exam_type = normalize_exam_type(exam_type)
    owned_exam(db, exam_type, exam_id, teacher_id)
    rubrics = db.query(ExamRubric).filter(
        ExamRubric.teacher_id == teacher_id,
        ExamRubric.exam_type == exam_type,
        _rubric_exam_filter(exam_type, exam_id),
    ).all()
    if not rubrics:
        raise HTTPException(404, "Rubric not found")
    for rubric in rubrics:
        if rubric.storage_path:
            Path(rubric.storage_path).unlink(missing_ok=True)
        db.delete(rubric)
    db.flush()
